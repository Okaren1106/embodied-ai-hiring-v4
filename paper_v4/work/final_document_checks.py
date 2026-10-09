"""Check delivered Word content, source hashes and editable objects after layout."""
from pathlib import Path
import hashlib,json,re,zipfile,xml.etree.ElementTree as ET
from docx import Document
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'release'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'm':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
checks=[]
def test(label,condition):
    checks.append({'check':label,'pass':bool(condition)})
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def clean(s):
    return s.replace('**','').replace('*','').replace('`','').replace('<br>','\n')
def source_tables(source):
    tables=[]
    for block in re.split(r'\n\s*\n',source):
        if not block.startswith('|'):continue
        rows=[]
        for line in block.splitlines():
            cells=[v.strip() for v in line.strip().strip('|').split('|')]
            if all(re.fullmatch(r':?-+:?',v) for v in cells):continue
            rows.append([clean(v) for v in cells])
        tables.append(rows)
    return tables
source={n:(ROOT/f'work/sol/Research_Working_{n}.md').read_text(encoding='utf-8') for n in ['Manuscript','Supplement']}
build=json.loads((ROOT/'verification/document_build.json').read_text())
for rel,digest in build['source_hashes'].items():test('locked source '+rel,sha(ROOT/rel)==digest)
for name,s in source.items():
    exported=(OUT/f'Research_Working_{name}.md').read_text(encoding='utf-8')
    test('Markdown release identical '+name,exported==s.replace('../../release/figures/','figures/'))
stats=json.loads((ROOT/'verification/word_statistics.json').read_text(encoding='utf-8-sig'))
objects={};docs={}
for name,table_kind,images in [('IJM_Research_Reading_Copy','Manuscript',2),('IJM_Main_Working_Text',None,0),
    ('IJM_Research_Tables','Manuscript',0),('IJM_Research_Supplement','Supplement',0)]:
    path=OUT/(name+'.docx');doc=Document(path);docs[name]=doc
    with zipfile.ZipFile(path) as z:
        xml=ET.fromstring(z.read('word/document.xml'))
        test(name+' no tracked changes',not xml.findall('.//w:ins',NS) and not xml.findall('.//w:del',NS))
        test(name+' no comments',not any('comments' in n for n in z.namelist()))
    expected=source_tables(source[table_kind]) if table_kind else []
    actual=[[[c.text for c in r.cells] for r in t.rows] for t in doc.tables]
    test(name+' editable tables equal source cells',actual==expected)
    test(name+' image count',len(doc.inline_shapes)==images)
    test(name+' anonymous author metadata',not doc.core_properties.author and not doc.core_properties.last_modified_by)
    if name!='IJM_Research_Tables':
        source_kind='Supplement' if name.endswith('Supplement') else 'Manuscript'
        math_count=len(re.findall(r'\\\(|\\\[',source[source_kind]))
        test(name+' all equations retained',len(xml.findall('.//m:oMath',NS))==math_count)
    if name in ['IJM_Main_Working_Text','IJM_Research_Reading_Copy']:
        powers=[ ''.join(x.itertext()) for x in xml.findall('.//m:sSup',NS)]
        test(name+' numerical exponent native math',any(x=='27' for x in powers))
    objects[name]={'sha256':sha(path),'tables':len(doc.tables),'images':len(doc.inline_shapes),
                   'equations':len(xml.findall('.//m:oMath',NS)),**stats[name]}
# Normal prose paragraphs should survive the formatter verbatim. Equations,
# lists, captions and tables have separate checks above or visual acceptance.
prose_count=0
for kind,name in [('Manuscript','IJM_Research_Reading_Copy'),('Supplement','IJM_Research_Supplement')]:
    paragraphs={p.text for p in docs[name].paragraphs}
    for block in re.split(r'\n\s*\n',source[kind]):
        b=block.strip()
        if not b or b.startswith(('#','|','![','- ',r'\[','**Human','**Table','**Supplementary Table','**Figure')):continue
        if b.startswith(('This is an internal','This supplement accompanies')) or r'\(' in b:continue
        if re.match(r'^\d+\. ',b):continue
        prose_count+=1;test(kind+' paragraph '+str(prose_count),clean(b.removeprefix('> ')) in paragraphs)
abstract=source['Manuscript'].split('## Structured abstract',1)[1].split('## 1.',1)[0]
abstract_words=len(abstract.split())
test('abstract including keywords/classification under 250',abstract_words<=250)
test('main 3 tables and 2 figures',objects['IJM_Research_Reading_Copy']['tables']==3 and objects['IJM_Research_Reading_Copy']['images']==2)
word_equivalent=stats['IJM_Main_Working_Text']['words']+5*280
test('IJM main plus 5 object allowance within 6000–8000',6000<=word_equivalent<=8000)
report={'pass':all(c['pass'] for c in checks),'checks':len(checks),'prose_paragraphs':prose_count,
    'structured_abstract_words_including_keywords_classification':abstract_words,
    'main_Word_words_plus_280_per_main_object':word_equivalent,'objects':objects,
    'results':checks,'limitations':'Checks verify content transfer and object integrity, not semantic accuracy or publisher permission.'}
(ROOT/'verification/document_content_checks.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k not in ['objects','results']}))
if not report['pass']:print(json.dumps([c for c in checks if not c['pass']],ensure_ascii=False))
assert report['pass']
