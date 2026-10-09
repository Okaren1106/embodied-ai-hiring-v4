"""Build editable Word reading, main-text, table and supplementary files.

Consumes the Sol-authored working manuscripts; no numerical claims are generated
by this formatter. Numerical tables remain editable. Equations use native OMML.
Run with the Codex bundled Python runtime (python-docx).
"""
from pathlib import Path
import re,json,hashlib,shutil
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'release'

def make(title,kind):
    d=Document();s=d.sections[0]
    s.page_width=Cm(21.59);s.page_height=Cm(27.94)
    s.top_margin=s.bottom_margin=Cm(2);s.left_margin=s.right_margin=Cm(2.1)
    s.header_distance=s.footer_distance=Cm(.85)
    for st in d.styles:
        for b in st.element.xpath('.//w:pBdr'):b.getparent().remove(b)
    for name,size in [('Normal',11),('Title',17),('Subtitle',10),('Heading 1',12),('Heading 2',11),('Caption',10),('Footer',9)]:
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
        fonts=st.element.get_or_add_rPr().rFonts
        for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:fonts.attrib.pop(qn('w:'+attr),None)
        fonts.set(qn('w:eastAsia'),'SimSun')
    pf=d.styles['Normal'].paragraph_format
    pf.line_spacing=1.35 if kind=='main' else 1.2;pf.space_after=Pt(6);pf.widow_control=True
    for name in ['Heading 1','Heading 2']:
        st=d.styles[name];st.font.bold=name=='Heading 1';st.font.italic=name=='Heading 2'
        st.paragraph_format.keep_with_next=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
    d.styles['Caption'].font.bold=True;d.styles['Caption'].font.italic=False
    d.styles['Caption'].paragraph_format.keep_with_next=True
    d.core_properties.title=title;d.core_properties.author='';d.core_properties.last_modified_by=''
    d.core_properties.subject='Research working manuscript';d.core_properties.comments=''
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
    return d

def mr(s):
    node=OxmlElement('m:r');t=OxmlElement('m:t');t.text=s;node.append(t);return node

def math_nodes(source):
    for a,b in [(r'\subseteq','⊆'),(r'\neq','≠'),(r'\Delta','Δ'),(r'\Pr','Pr'),
                (r'\qquad','    '),(r'\quad','  '),(r'\times','×'),(r'\cup','∪'),
                (r'\geq','≥'),(r'\leq','≤'),(r'\in','∈'),(r'\sum','Σ'),
                (r'\leftrightarrow','↔'),(r'\rightarrow','→'),(r'\{','{'),(r'\}','}'),(r'\,',' ')]:source=source.replace(a,b)
    assert '\\' not in source,'Unsupported equation command: '+source
    nodes=[]
    pattern=r'([A-Za-zΔ0-9]+)((?:[_^](?:\{[^}]+\}|[A-Za-z0-9]))*)|([^A-Za-zΔ0-9]+)'
    for m in re.finditer(pattern,source):
        if m[3]:nodes.append(mr(m[3]));continue
        base,scripts=m[1],m[2]
        attrs={op:value.strip('{}') for op,value in re.findall(r'([_^])(\{[^}]+\}|[A-Za-z0-9])',scripts)}
        if not attrs:nodes.append(mr(base));continue
        node=OxmlElement('m:sSubSup' if len(attrs)==2 else 'm:sSub' if '_' in attrs else 'm:sSup')
        e=OxmlElement('m:e');e.append(mr(base));node.append(e)
        for op,name in [('_','sub'),('^','sup')]:
            if op in attrs:
                ch=OxmlElement('m:'+name);ch.append(mr(attrs[op]));node.append(ch)
        nodes.append(node)
    return nodes

def math(p,source,display=False):
    om=OxmlElement('m:oMath')
    for n in math_nodes(source):om.append(n)
    if display:
        outer=OxmlElement('m:oMathPara');outer.append(om);p._p.append(outer)
        p.paragraph_format.keep_together=True
    else:p._p.append(om)

def runs(p,text):
    text=text.replace('<br>','\n')
    for token in re.split(r'(\\\(.*?\\\)|\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)',text):
        if not token:continue
        if token.startswith(r'\('):math(p,token[2:-2])
        elif token.startswith('**'):p.add_run(token[2:-2]).bold=True
        elif token.startswith('*'):p.add_run(token[1:-1]).italic=True
        elif token.startswith('`'):
            rr=p.add_run(token[1:-1]);rr.font.name='Consolas';rr.font.size=Pt(9)
        else:p.add_run(token)

def heading(text):
    return re.sub(r'[^\w\s]',' ',text,flags=re.UNICODE).strip()

def add_table(d,block):
    lines=block.splitlines();rows=[]
    for line in lines:
        row=[v.strip() for v in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r':?-+:?',v) for v in row):continue
        rows.append(row)
    n=len(rows[0]);assert all(len(r)==n for r in rows)
    # Data-specific widths favor labels while keeping numeric fields legible.
    widths=[3.8]+[(17.39-3.8)/(n-1)]*(n-1) if n>1 else [17.39]
    if n==6:widths=[3.2,3.2,1.2,1.2,2.0,6.59]
    if n==6 and rows[0][3]=='Firms':widths=[3.2,3.2,1.1,1.5,4.4,3.99]
    if n==8:widths=[2.1,1.7,1.0,1.0,1.0,3.0,3.0,4.59]
    if n==9:widths=[1.8,1.1,2.05,2.05,2.25,2.25,2.15,1.87,1.87]
    if n>=10:widths=[3.3]+[(17.39-3.3)/(n-1)]*(n-1)
    t=d.add_table(rows=1,cols=n);t.autofit=False
    for col,width in zip(t.columns,widths):col.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','bottom','insideH']:
        b=OxmlElement('w:'+name);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'B0B0B0');borders.append(b)
    t._tbl.tblPr.append(borders)
    for i,rr in enumerate(rows):
        row=t.rows[0] if i==0 else t.add_row();pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for j,(cell,val) in enumerate(zip(row.cells,rr)):
            cell.width=Cm(widths[j]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            p=cell.paragraphs[0];p.paragraph_format.line_spacing=1.05;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            p.paragraph_format.keep_with_next=i==0
            runs(p,val)
            for r in p.runs:r.font.size=Pt(9 if n>=8 else 9.5);r.bold=i==0
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)

def document(source,name,kind):
    blocks=[b.strip() for b in re.split(r'\n\s*\n',source) if b.strip()]
    title=next(b[2:] for b in blocks if b.startswith('# ') and not b.startswith('# INTERNAL'))
    d=make(title,kind);table_blocks=[];captions=[];i=0;references=False
    d.add_paragraph('AI assisted research working draft','Subtitle')
    d.add_paragraph('Prepared for author review. Publisher policy and author declarations remain the author’s responsibility.').runs[0].font.size=Pt(9)
    while i<len(blocks):
        b=blocks[i]
        if b.startswith('# INTERNAL') or b.startswith('This is an internal') or b.startswith('This supplement accompanies'):
            i+=1;continue
        if b.startswith('**Human author'):i+=1;continue
        if b.startswith('**Table ') or b.startswith('**Supplementary Table '):
            group=[b];j=i+1
            if j<len(blocks) and blocks[j].startswith('|'):group.append(blocks[j]);j+=1
            if j<len(blocks) and blocks[j].startswith('Notes:'):group.append(blocks[j]);j+=1
            assert len(group)>=2,b
            table_blocks.append(group)
            if kind=='main':d.add_paragraph('[Insert '+b.strip('*').split('.')[0]+' here]')
            else:
                d.add_paragraph(b.strip('*'),'Caption');add_table(d,group[1])
                if len(group)>2:
                    p=d.add_paragraph();runs(p,group[2]);p.paragraph_format.line_spacing=1.05
                    for r in p.runs:r.font.size=Pt(9)
            i=j;continue
        if b.startswith('**Figure '):
            imageblock=blocks[i+1];match=re.search(r'!\[(.*?)\]\((.*?)\)',imageblock);assert match,imageblock
            note=blocks[i+2] if i+2<len(blocks) and blocks[i+2].startswith('Notes:') else ''
            captions.append((b.strip('*'),note))
            if kind=='main':d.add_paragraph('[Insert '+b.strip('*').split('.')[0]+' here]')
            else:
                d.add_paragraph(b.strip('*'),'Caption');p=d.add_paragraph();p.paragraph_format.keep_with_next=True
                p.add_run().add_picture(str(OUT/'figures'/Path(match[2]).name),width=Cm(17.1))
                if note:
                    p=d.add_paragraph();runs(p,note)
                    for r in p.runs:r.font.size=Pt(9)
            i+=3 if note else 2;continue
        if b.startswith(r'\['):math(d.add_paragraph(),b[2:-2].strip(),True)
        elif b.startswith('# '):d.add_paragraph(heading(b[2:]),'Title')
        elif b.startswith('## '):
            d.add_paragraph(heading(b[3:]),'Heading 1');references=b[3:].strip()=='References'
        elif b.startswith('### '):d.add_paragraph(heading(b[4:]),'Heading 2')
        elif b.startswith('|'):add_table(d,b)
        elif b.startswith('- ') or re.match(r'^\d+\. ',b):
            for line in b.splitlines():
                p=d.add_paragraph(style='List Bullet' if line.startswith('- ') else 'List Number')
                runs(p,re.sub(r'^(?:- |\d+\. )','',line))
        else:
            p=d.add_paragraph();runs(p,b.removeprefix('> '))
            if references:
                p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5);p.paragraph_format.keep_together=True
            if b.startswith('Notes:'):
                for r in p.runs:r.font.size=Pt(9)
        i+=1
    if kind=='main':
        d.add_paragraph('Figure captions','Heading 1')
        for title,note in captions:
            d.add_paragraph(title,'Caption');p=d.add_paragraph();runs(p,note)
    d.save(OUT/(name+'.docx'))
    return table_blocks

def main():
    OUT.mkdir(exist_ok=True)
    main_path=ROOT/'work/sol/Research_Working_Manuscript.md';supp_path=ROOT/'work/sol/Research_Working_Supplement.md'
    main=main_path.read_text(encoding='utf-8');supp=supp_path.read_text(encoding='utf-8')
    tables=document(main,'IJM_Research_Reading_Copy','reader')
    document(main,'IJM_Main_Working_Text','main')
    document(supp,'IJM_Research_Supplement','supplement')
    d=make('Tables for technology qualification research','tables')
    d.add_paragraph('Tables for technology qualification research','Title')
    for i,group in enumerate(tables):
        if i:d.add_page_break()
        d.add_paragraph(group[0].strip('*'),'Caption');add_table(d,group[1])
        if len(group)>2:
            p=d.add_paragraph();runs(p,group[2])
            for r in p.runs:r.font.size=Pt(9)
    d.save(OUT/'IJM_Research_Tables.docx')
    for p in [main_path,supp_path]:
        text=p.read_text(encoding='utf-8').replace('../../release/figures/','figures/')
        (OUT/p.name).write_text(text,encoding='utf-8')
    manifest={'source_hashes':{str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [main_path,supp_path]},
              'outputs':[p.name for p in sorted(OUT.glob('*.docx'))],'main_tables':len(tables),
              'render_and_visual_inspection_required':True}
    (ROOT/'verification/document_build.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(manifest))

if __name__=='__main__':main()
