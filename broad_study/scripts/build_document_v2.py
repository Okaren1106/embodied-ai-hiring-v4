"""Markdown to Word. Authors: Astra (prior) + grok4.6 (this package)."""
from pathlib import Path
import re,json,zipfile,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'document/具身智能公开招聘研究_阶段E综合报告_v2_grok4.6.docx'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.78)
def font(style,size,bold=False):
 s=doc.styles[style];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.bold=bold
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if style=='Normal' else '微软雅黑')
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.18
font('Normal',11);font('Title',22,True)
for h,sz in [('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
 font(h,sz,True);doc.styles[h].paragraph_format.space_before=Pt(14)
doc.core_properties.author='grok4.6';doc.core_properties.title='具身智能公开招聘研究综合报告 v2'
def inline(p,text):
 pattern=r'(\*\*.*?\*\*)'
 for token in re.split(pattern,text):
  if not token:continue
  if token.startswith('**'):p.add_run(token[2:-2]).bold=True
  else:p.add_run(token)
def make_table(rows):
 rows=[r for r in rows if not re.match(r'^\|[\-: |]+\|$',r)]
 vals=[[x.strip() for x in r.strip('|').split('|')] for r in rows];n=len(vals[0]);tb=doc.add_table(rows=0,cols=n)
 for i,rr in enumerate(vals):
  row=tb.add_row()
  for j,(cell,txt) in enumerate(zip(row.cells,rr)):
   p=cell.paragraphs[0];inline(p,txt)
   for run in p.runs:run.font.size=Pt(9);run.bold=i==0
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
# concatenate main report + stage E
text=(ROOT/'REPORT.md').read_text(encoding='utf-8')+'\n\n'+(ROOT/'reports/阶段E初稿.md').read_text(encoding='utf-8')
lines=text.splitlines();i=0
while i<len(lines):
 line=lines[i];i+=1
 if not line.strip():continue
 if line.startswith('|'):
  rows=[line]
  while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
  make_table(rows);continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
  img=ROOT/m.group(2)
  bmp=img.with_suffix('.bmp')
  use=bmp if bmp.exists() else (img if img.exists() and img.suffix.lower() in ('.png','.jpg','.jpeg','.bmp','.gif') else None)
  if use is not None:
   p=doc.add_paragraph();run=p.add_run();run.add_picture(str(use),width=Inches(6.4))
   cap=doc.add_paragraph(m.group(1))
   for r in cap.runs:r.font.size=Pt(9)
 elif line.startswith('# '):inline(doc.add_paragraph(style='Title'),line[2:])
 elif line.startswith('## '):inline(doc.add_paragraph(style='Heading 1'),line[3:])
 elif line.startswith('### '):inline(doc.add_paragraph(style='Heading 2'),line[4:])
 else:
  p=doc.add_paragraph();inline(p,line)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Astra前期 + grok4.6接续  |  ').font.size=Pt(9)
OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
pub=Path(r'D:/Embod AI analysis/embodied-ai-hiring-research/output/doc')
pub.mkdir(parents=True,exist_ok=True)
shutil.copy2(OUT, pub/OUT.name)
with zipfile.ZipFile(OUT) as z: assert z.testzip() is None
qa={'docx':OUT.name,'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'tables':len(doc.tables),'images':len(doc.inline_shapes),'bytes':OUT.stat().st_size}
(ROOT/'verification/document_structure.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(qa,ensure_ascii=False))
