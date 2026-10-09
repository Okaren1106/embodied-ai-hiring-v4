from pathlib import Path
import re,json,zipfile,hashlib
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'document/具身智能公开招聘研究_阶段D独立研究报告_Astra_v1.docx'
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.78)
sec.footer_distance=Inches(.32)
def font(style,size,bold=False):
 s=doc.styles[style];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.bold=bold
 s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体' if style=='Normal' else '微软雅黑')
 s.paragraph_format.space_after=Pt(6);s.paragraph_format.line_spacing=1.18
font('Normal',11);font('Title',25,True);font('Subtitle',16,True)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Normal'].paragraph_format.line_spacing=1.15
doc.styles['Subtitle'].font.italic=False
for s in doc.styles:
 for border in list(s.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for h,sz in [('Heading 1',16),('Heading 2',12.5),('Heading 3',11.5)]:
 font(h,sz,True);doc.styles[h].paragraph_format.space_before=Pt(14);doc.styles[h].paragraph_format.keep_with_next=True
doc.core_properties.author='Astra';doc.core_properties.title='具身智能公开招聘研究：阶段D独立研究报告（Astra）';doc.core_properties.subject='公司工作组合、共同职能与交付责任';doc.core_properties.comments=''
def inline(p,text):
 pattern=r'(\*\*.*?\*\*|\[[^\]]+\]\([^\)]+\))'
 for token in re.split(pattern,text):
  if not token:continue
  if token.startswith('**'):p.add_run(token[2:-2]).bold=True
  elif token.startswith('[') and '](' in token:
   m=re.match(r'\[([^\]]+)\]\(([^\)]+)\)',token);label,url=m.groups();h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'285A83');pr.append(color);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h)
  else:p.add_run(token)
def make_table(rows):
 rows=[r for r in rows if not re.match(r'^\|[\-: |]+\|$',r)]
 vals=[[x.strip() for x in r.strip('|').split('|')] for r in rows];n=len(vals[0]);tb=doc.add_table(rows=0,cols=n);tb.autofit=False
 # Compact value columns and generous interpretation columns.
 widths={3:[1.6,1.75,3.59],4:[1.6,1.76,1.8,1.78],6:[1.15,.85,.85,.85,1.2,2.04]}.get(n,[6.94/n]*n)
 for c,w in zip(tb.columns,widths):c.width=Inches(w)
 for i,rr in enumerate(vals):
  row=tb.add_row();trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
  if i==0:rep=OxmlElement('w:tblHeader');trpr.append(rep)
  for j,(cell,txt) in enumerate(zip(row.cells,rr)):
   cell.width=Inches(widths[j]);p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.1;p.paragraph_format.keep_with_next=i<len(vals)-1
   inline(p,txt)
   for run in p.runs:run.font.size=Pt(9.5);run.bold=i==0
   tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');marg.append(el)
   tcpr.append(marg)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D5D8DB');borders.append(el)
   tcpr.append(borders)
   if i==0:shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'EEF0F2');tcpr.append(shd)
 doc.add_paragraph().paragraph_format.space_after=Pt(0)
lines=(ROOT/'REPORT.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
 line=lines[i];i+=1
 if not line.strip():continue
 if line.startswith('|'):
  rows=[line]
  while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
  make_table(rows);continue
 if line.startswith('!['):
  m=re.match(r'!\[(.*?)\]\((.*?)\)',line);p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;run=p.add_run();shape=run.add_picture(str(ROOT/m.group(2)),width=Inches(6.86));shape._inline.docPr.set('descr',m.group(1));p=doc.add_paragraph(m.group(1));p.paragraph_format.space_after=Pt(7)
  for r in p.runs:r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('444444')
 elif line.startswith('# '):inline(doc.add_paragraph(style='Title'),line[2:])
 elif line.startswith('## 阶段'):inline(doc.add_paragraph(style='Subtitle'),line[3:])
 elif line.startswith('### '):inline(doc.add_paragraph(style='Heading 2'),line[4:])
 elif line.startswith('## '):
  p=doc.add_paragraph(style='Heading 1');inline(p,line[3:])
  if line.startswith('## 12.'):p.paragraph_format.page_break_before=True
 else:
  p=doc.add_paragraph();inline(p,line);p.paragraph_format.widow_control=True
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Astra  |  ');r.font.size=Pt(9);r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
OUT.parent.mkdir(exist_ok=True);doc.save(OUT)
with zipfile.ZipFile(OUT) as z:
 assert z.testzip() is None
 xml=z.read('word/document.xml').decode();assert 'turn6view' not in xml and 'Lorem ipsum' not in xml
 links=len([x for x in doc.part.rels.values() if x.reltype==RT.HYPERLINK]);assert links==8,links
qa={'docx':OUT.name,'sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),'tables':len(doc.tables),'images':len(doc.inline_shapes),'external_links':links,'structural_check':'PASS','visual_review':'PENDING'}
(ROOT/'verification/document_structure.json').write_text(json.dumps(qa,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(qa,ensure_ascii=False))
